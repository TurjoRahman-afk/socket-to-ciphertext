"""Turn the project's markdown documents into Word files for submission.

    python -m tools.to_word

Writes .docx into docs-word/, one per document, overwriting what is there.
The markdown stays the source of truth: edit `docs/*.md`, run this again, and
the Word files catch up. Editing the .docx directly works too, but the next
run will overwrite it -- so do that only once, on the copy you are about to
hand in.

Not part of the application. `python-docx` is a tool dependency and is listed
in requirements-dev.txt rather than requirements.txt, so nobody has to install
it to run Semaphore.

The converter handles the subset of markdown these documents actually use:
headings, paragraphs, bullet and numbered lists, tables, fenced code blocks,
block quotes, horizontal rules, and inline bold/italic/code. It is not a
general markdown implementation and does not pretend to be.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError:  # pragma: no cover -- the tool is optional
    sys.exit("python-docx is needed for this:  pip install -r requirements-dev.txt")

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs"
OUT = ROOT / "docs-word"

#: Which documents to convert, and what to call the file a marker opens.
#: The submission ones come first because that is the order they are read in.
DOCUMENTS = [
    ("design.md", "Semaphore - Design Report.docx"),
    ("testing-report.md", "Semaphore - Testing Report.docx"),
    ("reflection-template.md", "Semaphore - Reflection Template.docx"),
    ("threat-model.md", "Semaphore - Threat Model.docx"),
    ("protocol.md", "Semaphore - Protocol Specification.docx"),
    ("demo.md", "Semaphore - Demonstration Runbook.docx"),
    ("submission.md", "Semaphore - Submission Checklist.docx"),
]

INK = RGBColor(0x2B, 0x2B, 0x2B)
MUTED = RGBColor(0x60, 0x60, 0x60)
ACCENT = RGBColor(0xC0, 0x55, 0x1F)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text: str, *, base_bold: bool = False) -> None:
    """Write text into a paragraph, honouring **bold**, *italic* and `code`."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = ACCENT
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(piece)
        if base_bold:
            run.bold = True


def strip_links(text: str) -> str:
    """[label](target) -> label. A Word reader cannot follow a repo path."""
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_divider(line: str) -> bool:
    """The |---|---| line under a table header."""
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def add_table(document, lines: list[str]) -> None:
    rows = [split_row(line) for line in lines if not is_divider(line)]
    if not rows:
        return
    width = max(len(row) for row in rows)

    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, strip_links(row[c] if c < len(row) else ""), base_bold=r == 0)
            for run in paragraph.runs:
                run.font.size = Pt(9)
    document.add_paragraph()


def add_code(document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = INK


def convert(source: Path, destination: Path) -> None:
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    lines = source.read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Fenced code
        if stripped.startswith("```"):
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(document, block)
            i += 1
            continue

        # Table: a pipe row followed by a divider
        if stripped.startswith("|") and i + 1 < len(lines) and is_divider(lines[i + 1]):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(document, block)
            continue

        if not stripped:
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            rule = document.add_paragraph()
            rule.paragraph_format.space_before = Pt(2)
            run = rule.add_run("_" * 62)
            run.font.color.rgb = RGBColor(0xC8, 0xC8, 0xC8)
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = strip_links(stripped[level:].strip())
            heading = document.add_heading(level=min(level, 4))
            heading.text = ""
            add_inline(heading, text)
            i += 1
            continue

        if stripped.startswith(">"):
            quote = document.add_paragraph()
            quote.paragraph_format.left_indent = Pt(24)
            add_inline(quote, strip_links(stripped.lstrip("> ").strip()))
            for run in quote.runs:
                run.italic = True
                run.font.color.rgb = MUTED
            i += 1
            continue

        if re.match(r"^[-*] ", stripped):
            bullet = document.add_paragraph(style="List Bullet")
            add_inline(bullet, strip_links(stripped[2:]))
            i += 1
            continue

        if re.match(r"^\d+\. ", stripped):
            number = document.add_paragraph(style="List Number")
            add_inline(number, strip_links(re.sub(r"^\d+\.\s*", "", stripped)))
            i += 1
            continue

        body = document.add_paragraph()
        body.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(body, strip_links(stripped))
        i += 1

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> int:
    OUT.mkdir(exist_ok=True)
    written = 0
    for name, title in DOCUMENTS:
        source = SOURCE / name
        if not source.exists():
            print(f"  skipped {name} (not written yet)")
            continue
        convert(source, OUT / title)
        print(f"  {title}")
        written += 1

    # Anybody's reflection, once they have written one.
    for reflection in sorted(SOURCE.glob("reflection-*.md")):
        if reflection.name == "reflection-template.md":
            continue
        who = reflection.stem.replace("reflection-", "").title()
        convert(reflection, OUT / f"Semaphore - Reflection - {who}.docx")
        print(f"  Semaphore - Reflection - {who}.docx")
        written += 1

    print(f"\n{written} documents in {OUT.relative_to(ROOT)}/")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
